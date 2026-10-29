from __future__ import annotations

from pathlib import Path


def _created(props) -> dict:
    """{"created": "YYYY-MM-DD"} from the document properties, if plausible."""
    d = getattr(props, "created", None)
    return {"created": d.date().isoformat()} if d is not None and d.year >= 1990 else {}


def extract(path: Path, max_chars: int) -> tuple[str, dict]:
    ext = path.suffix.lower()
    if ext == ".docx":
        return _docx(path, max_chars)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(path, max_chars)
    if ext == ".pptx":
        return _pptx(path, max_chars)
    return "", {}  # legacy/OpenDocument formats: classified by name and type only


def _docx(path: Path, max_chars: int) -> tuple[str, dict]:
    import docx

    d = docx.Document(str(path))
    meta = _created(d.core_properties)
    if d.core_properties.title:
        meta["title"] = d.core_properties.title
    out, total = [], 0
    for p in d.paragraphs:
        if p.text.strip():
            out.append(p.text)
            total += len(p.text)
            if total >= max_chars:
                break
    return " ".join(out), meta


def _xlsx(path: Path, max_chars: int) -> tuple[str, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    meta = {"sheets": wb.sheetnames[:10], **_created(wb.properties)}
    out, total = [], 0
    for ws in wb.worksheets[:3]:
        for row in ws.iter_rows(max_row=20, values_only=True):
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                line = " | ".join(cells)
                out.append(line)
                total += len(line)
            if total >= max_chars:
                break
    wb.close()
    return " ; ".join(out), meta


def _pptx(path: Path, max_chars: int) -> tuple[str, dict]:
    from pptx import Presentation

    prs = Presentation(str(path))
    meta = {"slides": len(prs.slides), **_created(prs.core_properties)}
    out, total = [], 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
                total += len(shape.text_frame.text)
        if total >= max_chars:
            break
    return " ".join(out), meta
